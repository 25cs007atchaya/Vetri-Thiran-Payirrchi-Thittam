import os

from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH

from fpdf import FPDF


# =========================================================
# TEXT CLEANING
# =========================================================

def sanitize_text(text: str) -> str:
    """
    Clean generated text before formatting or exporting.
    """

    if not text:
        return ""

    text = text.replace("\r\n", "\n")
    text = text.replace("\r", "\n")

    return text.strip()


# =========================================================
# HTML PREVIEW
# =========================================================

def format_html_preview(text: str) -> str:
    """
    Convert generated document text into an HTML preview.
    """

    text = sanitize_text(text)

    paragraphs = text.split("\n")

    html_content = ""

    for paragraph in paragraphs:

        paragraph = paragraph.strip()

        if paragraph:
            html_content += f"<p>{paragraph}</p>"

    return f"""
    <div class="legal-document">
        {html_content}
    </div>
    """


# =========================================================
# DOCX FORMATTER
# =========================================================

def format_docx(text: str, output_path: str):
    """
    Create a professionally formatted DOCX document.
    """

    text = sanitize_text(text)

    document = Document()

    # Page settings
    section = document.sections[0]

    section.top_margin = Inches(0.7)
    section.bottom_margin = Inches(0.7)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)

    # Default font
    style = document.styles["Normal"]

    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    # Header
    header = section.header

    header_paragraph = header.paragraphs[0]

    header_paragraph.text = "LEGALEASE"
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    header_run = header_paragraph.runs[0]

    header_run.bold = True
    header_run.font.name = "Times New Roman"
    header_run.font.size = Pt(16)

    # Document content
    lines = text.split("\n")

    first_content = True

    for line in lines:

        line = line.strip()

        if not line:
            continue

        # Main title
        if first_content:

            paragraph = document.add_paragraph()

            paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

            run = paragraph.add_run(line)

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(18)

            paragraph.paragraph_format.space_after = Pt(14)

            first_content = False

        # Headings
        elif (
            line.isupper()
            or line.endswith(":")
            or line.startswith("1.")
            or line.startswith("2.")
            or line.startswith("3.")
            or line.startswith("4.")
            or line.startswith("5.")
        ):

            paragraph = document.add_paragraph()

            run = paragraph.add_run(line)

            run.bold = True
            run.font.name = "Times New Roman"
            run.font.size = Pt(13)

            paragraph.paragraph_format.space_before = Pt(8)
            paragraph.paragraph_format.space_after = Pt(5)

        # Normal paragraph
        else:

            paragraph = document.add_paragraph()

            run = paragraph.add_run(line)

            run.font.name = "Times New Roman"
            run.font.size = Pt(12)

            paragraph.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY

            paragraph.paragraph_format.space_after = Pt(6)

    # Footer
    footer = section.footer

    footer_paragraph = footer.paragraphs[0]

    footer_paragraph.text = (
        "Generated using LegalEase – AI-Powered Legal Document Generator"
    )

    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER

    footer_run = footer_paragraph.runs[0]

    footer_run.font.name = "Times New Roman"
    footer_run.font.size = Pt(9)

    document.save(output_path)


# =========================================================
# PDF FORMATTER
# =========================================================

def format_pdf(text: str, output_path: str):
    """
    Create a professionally formatted Unicode PDF document.
    """

    text = sanitize_text(text)

    # Get the LegalEase project folder
    project_root = os.path.dirname(
        os.path.dirname(os.path.abspath(__file__))
    )

    # Font paths
    regular_font = os.path.join(
        project_root,
        "fonts",
        "NotoSans-Regular.ttf"
    )

    bold_font = os.path.join(
        project_root,
        "fonts",
        "NotoSans-Bold.ttf"
    )

    # Check fonts
    if not os.path.exists(regular_font):
        raise FileNotFoundError(
            f"Font not found: {regular_font}"
        )

    if not os.path.exists(bold_font):
        raise FileNotFoundError(
            f"Font not found: {bold_font}"
        )

    # Create PDF
    pdf = FPDF()

    pdf.set_auto_page_break(
        auto=True,
        margin=20
    )

    pdf.add_page()

    # Add Unicode fonts
    pdf.add_font(
        "Noto",
        style="",
        fname=regular_font
    )

    pdf.add_font(
        "Noto",
        style="B",
        fname=bold_font
    )

    # -----------------------------------------------------
    # HEADER
    # -----------------------------------------------------

    pdf.set_font(
        "Noto",
        style="B",
        size=16
    )

    pdf.cell(
        0,
        10,
        "LEGALEASE",
        align="C"
    )

    pdf.ln(12)

    # -----------------------------------------------------
    # DOCUMENT CONTENT
    # -----------------------------------------------------

    lines = text.split("\n")

    first_content = True

    for line in lines:

        line = line.strip()

        if not line:
            continue

        # Main title
        if first_content:

            pdf.set_font(
                "Noto",
                style="B",
                size=18
            )

            pdf.multi_cell(
                0,
                10,
                line,
                align="C"
            )

            pdf.ln(5)

            first_content = False

        # Heading
        elif (
            line.isupper()
            or line.endswith(":")
            or line.startswith("1.")
            or line.startswith("2.")
            or line.startswith("3.")
            or line.startswith("4.")
            or line.startswith("5.")
        ):

            pdf.set_font(
                "Noto",
                style="B",
                size=13
            )

            pdf.multi_cell(
                0,
                8,
                line
            )

            pdf.ln(2)

        # Normal paragraph
        else:

            pdf.set_font(
                "Noto",
                style="",
                size=11
            )

            pdf.multi_cell(
                0,
                7,
                line
            )

            pdf.ln(3)

    # -----------------------------------------------------
    # FOOTER
    # -----------------------------------------------------

    total_pages = pdf.page_no()

    for page_number in range(1, total_pages + 1):

        pdf.page = page_number

        pdf.set_y(-15)

        pdf.set_font(
            "Noto",
            style="",
            size=8
        )

        pdf.cell(
            0,
            10,
            f"LegalEase | Page {page_number}",
            align="C"
        )

    # Save PDF
    pdf.output(output_path)